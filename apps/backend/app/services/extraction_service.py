"""
AI extraction service — extract topics from syllabus PDFs using Groq.

Pipeline:
1. Download PDF from R2
2. Extract text via pypdf
3. Send to Groq via the model router (circuit-breaker fallback)
4. Parse response into Topic objects
5. Bulk insert into DB, linked to the source document
"""

import io
import json
import logging
from datetime import datetime, timezone

from pypdf import PdfReader
from sqlalchemy.ext.asyncio import AsyncSession
from fastapi.concurrency import run_in_threadpool

from app.models.document import Document
from app.models.topic import Topic
from app.services import storage_service
from app.services.groq_router import get_router

logger = logging.getLogger(__name__)

# ──────────────────────────────────────────────────────────────────────────────
# System prompts
# ──────────────────────────────────────────────────────────────────────────────
EXTRACTION_SYSTEM_PROMPT_SYLLABUS = """You are an academic syllabus parser. Your job is to extract the list of course topics/modules from a syllabus PDF.

Rules:
1. Extract each distinct topic, module, chapter, or unit mentioned in the syllabus.
2. Maintain the order they appear in the syllabus.
3. Be concise — use the title as written. Don't add commentary.
4. If the syllabus has a weekly schedule, extract the topic for each week.
5. Ignore administrative items (grading policies, office hours, etc.) — focus on content topics only.

Respond ONLY with a JSON array of topic strings. Example:
["Introduction to Algorithms", "Sorting and Searching", "Graph Theory", "Dynamic Programming"]

If you cannot extract any topics, respond with an empty array: []"""

EXTRACTION_SYSTEM_PROMPT_SLIDES = """You are an academic lecture content parser. Your job is to extract the list of topics/concepts covered in lecture slides or instructor notes.

Rules:
1. Extract each distinct topic, concept, or learning objective from the slides/notes.
2. Maintain the order they appear (slide order for slides, document order for notes).
3. Be concise — use the title as written. Don't add commentary.
4. Ignore administrative slides (title slide, agenda, summary, references, office hours) — focus on content topics only.
5. If slides have clear section headers, use those as topics.

Respond ONLY with a JSON array of topic strings. Example:
["Introduction to Machine Learning", "Linear Regression", "Gradient Descent", "Overfitting and Regularization"]

If you cannot extract any topics, respond with an empty array: []"""

# Max characters sent to the LLM (keeps well within all models' context windows)
MAX_SYLLABUS_CHARS = 25_000


# ──────────────────────────────────────────────────────────────────────────────
# PDF helpers
# ──────────────────────────────────────────────────────────────────────────────

def extract_text_from_pdf(pdf_bytes: bytes) -> str:
    """Extract all text content from a PDF file."""
    reader = PdfReader(io.BytesIO(pdf_bytes))
    parts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            parts.append(text)
    return "\n\n".join(parts)


def extract_text_from_pptx(pptx_bytes: bytes) -> str:
    """Extract all text content from a PPTX file (lecture slides)."""
    from pptx import Presentation
    import io as _io
    prs = Presentation(_io.BytesIO(pptx_bytes))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        parts.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        text = cell.text.strip()
                        if text:
                            parts.append(text)
    return "\n\n".join(parts)


def _parse_topics_from_response(response_text: str) -> list[str]:
    """
    Parse a JSON array of topic strings out of the LLM response.
    Handles optional markdown code fences gracefully.
    """
    text = response_text.strip()

    # Strip markdown code fences if present
    if text.startswith("```"):
        lines = text.splitlines()
        text = "\n".join(lines[1:-1]).strip()

    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        logger.error("Failed to parse LLM response as JSON: %.200s", text)
        return []

    if not isinstance(parsed, list):
        logger.warning("LLM returned non-list JSON: %.200s", text)
        return []

    return [str(t).strip() for t in parsed if t]


# ──────────────────────────────────────────────────────────────────────────────
# Public API
# ──────────────────────────────────────────────────────────────────────────────

def call_groq_for_topics(text: str, system_prompt: str) -> tuple[list[str], str]:
    """
    Send text to Groq via the circuit-breaker router.

    Args:
        text: The text to extract topics from
        system_prompt: The system prompt to use (syllabus vs slides)

    Returns:
        (topics, model_used) — list of topic title strings and the model name.

    Raises:
        RuntimeError if all models are exhausted.
    """
    if len(text) > MAX_SYLLABUS_CHARS:
        text = text[:MAX_SYLLABUS_CHARS] + "\n\n[... truncated for length]"

    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": f"Extract the topics from this content:\n\n{text}"},
    ]

    router = get_router()
    response_text, model_used = router.chat(messages=messages, temperature=0.1, max_tokens=2000)

    topics = _parse_topics_from_response(response_text)
    return topics, model_used


async def extract_topics_for_document(
    document: Document,
    db: AsyncSession,
) -> list[Topic]:
    """
    Full extraction pipeline for a single document:
      1. Download file from R2
      2. Extract text via pypdf (PDF) or python-pptx (PPTX)
      3. Call Groq (with circuit-breaker fallback) for topics
      4. Insert Topic rows into the DB, linked to this document

    Updates `document.processing_status` and `document.error_message` in place.
    """
    document.processing_status = "processing"
    await db.flush()

    try:
        # Step 1: Download file from R2
        file_bytes = await run_in_threadpool(storage_service.download_file, document.r2_key)

        # Step 2: Extract text based on MIME type
        mime_type = (document.mime_type or "").lower()
        if mime_type == "application/pdf":
            text = await run_in_threadpool(extract_text_from_pdf, file_bytes)
            system_prompt = EXTRACTION_SYSTEM_PROMPT_SYLLABUS
        elif mime_type in ("application/vnd.openxmlformats-officedocument.presentationml.presentation",
                           "application/vnd.ms-powerpoint"):
            text = await run_in_threadpool(extract_text_from_pptx, file_bytes)
            system_prompt = EXTRACTION_SYSTEM_PROMPT_SLIDES
        else:
            document.processing_status = "failed"
            document.error_message = f"Unsupported file type for topic extraction: {mime_type}"
            await db.flush()
            return []

        if not text.strip():
            document.processing_status = "failed"
            document.error_message = "Could not extract any text from the document"
            await db.flush()
            return []

        # Step 3: Call Groq with fallback
        topic_titles, model_used = await run_in_threadpool(call_groq_for_topics, text, system_prompt)
        logger.info(
            "Extracted %d topics from document %d using model %s",
            len(topic_titles), document.id, model_used,
        )

        if not topic_titles:
            document.processing_status = "failed"
            document.error_message = "AI could not identify any topics in this document"
            await db.flush()
            return []

        # Step 4: Bulk insert topics
        created_topics: list[Topic] = []
        for idx, title in enumerate(topic_titles):
            topic = Topic(
                course_id=document.course_id,
                user_id=document.user_id,
                source_document_id=document.id,
                title=title,
                order_index=idx,
                is_confirmed=False,
            )
            db.add(topic)
            created_topics.append(topic)

        document.processing_status = "processed"
        document.processed_at = datetime.now(timezone.utc)
        document.error_message = None
        await db.flush()

        return created_topics

    except Exception as exc:
        logger.exception("Extraction failed for document %d", document.id)
        document.processing_status = "failed"
        document.error_message = str(exc)[:500]
        await db.flush()
        return []
